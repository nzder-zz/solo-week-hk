"""Duo-Week-Planner.xlsx: every price is an input cell (blue); totals, shares and fit are formulas.
Shape: together Sat 31 Oct -> Wed 4 Nov (she flies back to HK on Wed evening), you alone Thu-Fri, home Sat 7 Nov."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F)
GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True)
HEAD = Font(name=F, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="241626")
YELLOW = PatternFill("solid", fgColor="FFFF00")
EUR = '€#,##0;(€#,##0);-'
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS_TXT = {"gf27": "Live, Google Flights 27 Oct", "est": "Estimate (schedule checked 28 Oct)"}
CHK = {"ok": "Open / fine", "warn": "Open, with a caveat", "out": "Closed or damaged: left out", "todo": "Still to verify"}
WHO = {"both": "Together", "split": "She flies tonight", "you": "You alone"}


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.alignment = HEAD, HFILL, Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def put(ws, r, c, v, font=BLACK, fmt=None, wrap=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if wrap:
        cell.alignment = WRAP
    return cell


def build_xlsx(D, path):
    wb = Workbook()
    fares = D["fares"]
    P = D["params"]

    # ------------------------------------------------ Read me
    ws = wb.active
    ws.title = "Read me"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Week for two after Hong Kong — together Sat 31 Oct → Wed 4 Nov, then you alone until Sat 7 Nov", BOLD),
        ("Leaning to Japan: Tokyo + one smaller town (Kawaguchiko, Hakone or Nikko), with Kansai, Kyushu and Taiwan kept as other options. Each trip has one chill plan. Private rooms while together, taxis split, no Kyoto, no bike trips. She flies back to HK on Wed 4 Nov in the evening; you spend Thu 5 and Fri 6 on your own and fly home on Sat 7 Nov (single ticket).", BLACK),
        ("", BLACK),
        ("How to use", BOLD),
        ("• My plan: pick a trip from the drop-down; the costs and both shares update.", BLACK),
        ("• Overview: every trip side by side.", BLACK),
        ("• Change any blue number (Flights, Rooms, Plans, Inputs, Scores): every total is a formula.", BLACK),
        ("• Checks: what was verified on 28 Oct (Typhoon Champi damage, closures, volcano levels, fares) and what is still open.", BLACK),
        ("", BLACK),
        ("Colours", BOLD),
        ("Blue text = an input you can change · black = formula · green = link to another sheet · yellow fill = the choices on My plan.", BLACK),
        ("", BLACK),
        ("Price status", BOLD),
        ("Flights out (31 Oct) and your flight home (Sat 7 Nov): live fares from Google Flights, 27 Oct, for 1 adult (two seats can cost a little more).", BLACK),
        ("Her flight back on Wed 4 Nov: evening schedules checked 28 Oct, prices are estimates — no browser was reachable to check them live.", BLACK),
        ("Rooms: typical prices for the cheapest decent private double while together, and a hostel bed or single room for your nights alone. Not quotes.", BLACK),
        ("", BLACK),
        ("Shares: each of you pays your own flights; the days together are split 50/50; your days alone are yours.", BLACK),
        ("Rates used: 1 EUR = ¥177 · NT$36.5. Nothing is booked.", BLACK),
    ]
    for i, (t, f) in enumerate(lines, 1):
        put(ws, i, 1, t, f, wrap=True)

    # ------------------------------------------------ Inputs
    wi = wb.create_sheet("Inputs")
    header(wi, 1, ["Input", "Value", "Note"], [40, 12, 80])
    rows = [("Days of food together (for two)", P["days_together"], "Sat evening to Wed afternoon", "0"),
            ("Days of food on your own", P["days_solo"], "Wed evening to Sat", "0"),
            ("Misc per person (€)", P["misc_pp"], "SIM, tips, laundry (while together)", EUR),
            ("Meal credit (share of a day's food)", P["meal_credit"], "When dinner + breakfast come with the room (temple, minshuku, ryokan)", "0.00")]
    for i, (a, b, c, f) in enumerate(rows, 2):
        put(wi, i, 1, a); put(wi, i, 2, b, BLUE, f); put(wi, i, 3, c, wrap=True)
    DAYS, DSOLO, MISC, CREDIT = "Inputs!$B$2", "Inputs!$B$3", "Inputs!$B$4", "Inputs!$B$5"

    # ------------------------------------------------ Flights
    wf = wb.create_sheet("Flights")
    header(wf, 1, ["Key", "Leg", "Route", "Flight", "Price per person (€)", "Status", "Alternatives"], [16, 26, 13, 70, 14, 30, 70])
    r = 2
    groups = [("OUT", "Both, out Sat 31 Oct (each)", fares["out"]), ("HER", "Her, back to HK Wed 4 Nov", fares["her"]), ("HIS", "You, home Sat 7 Nov", fares["his"])]
    for pre, leg, dct in groups:
        for k, v in dct.items():
            for col, val in enumerate([f"{pre}|{k}", leg, v["route"], v["detail"], v["price"], STATUS_TXT[v["status"]], v.get("alt", "")], 1):
                put(wf, r, col, val, BLUE if col == 5 else BLACK, EUR if col == 5 else None, wrap=col in (4, 7))
            r += 1
    FL_LAST = r - 1
    FKEY, FPRICE = f"Flights!$A$2:$A${FL_LAST}", f"Flights!$E$2:$E${FL_LAST}"

    # ------------------------------------------------ Rooms
    wr = wb.create_sheet("Rooms")
    header(wr, 1, ["Plan key", "Trip", "Plan", "Night", "Who", "Where", "Room (€)", "Dinner + breakfast included"], [28, 40, 9, 12, 11, 46, 11, 14])
    r = 2
    for d in D["dests"]:
        for o in d["options"]:
            key = f"{d['id']}|{o['style']}"
            for i, n in enumerate(o["nights"]):
                for col, val in enumerate([key, d["title"], o["style"], D["dates"][i], "You" if n["solo"] else "Together", n["place"], n["room"], "Yes" if n["meals"] else "No"], 1):
                    put(wr, r, col, val, BLUE if col in (7, 8) else BLACK, EUR if col == 7 else None)
                r += 1
    RM_LAST = r - 1
    RKEY, RWHO, RPRICE, RMEAL = (f"Rooms!${c}$2:${c}${RM_LAST}" for c in "AEGH")
    put(wr, r + 1, 1, "Together: price of a private double for the two of you. You: a hostel bed or single room. Typical rates, not quotes (28 Oct).")

    # ------------------------------------------------ Plans (the cost engine)
    wp = wb.create_sheet("Plans")
    cols = ["Plan key", "Trip", "Plan", "Food / day / person (together)", "Rooms together", "Nights with meals (together)", "Transport / person (together)",
            "Shared taxis & cars", "Activities / person (together)", "Flight out / person", "Her flight to HK", "Your flight home", "Food together (two)",
            "Transport together", "Activities together", "Misc", "Together: total", "Your food / day (alone)", "Your rooms alone", "Your nights with meals",
            "Your food alone", "Your transport alone", "Your activities alone", "Your days alone: total", "Total for two", "Your share", "Her share",
            "Lookup", "Out key", "Her key", "Your key"]
    header(wp, 1, cols, [22, 36, 8] + [11] * 24 + [30, 12, 10, 10])
    r = 2
    plan_rows = {}
    for d in D["dests"]:
        out_k = [k for k, v in fares["out"].items() if v == d["out"]][0]
        her_k = [k for k, v in fares["her"].items() if v == d["her"]][0]
        his_k = [k for k, v in fares["his"].items() if v == d["his"]][0]
        for o in d["options"]:
            key = f"{d['id']}|{o['style']}"
            plan_rows[key] = r
            vals = {
                1: key, 2: d["title"], 3: o["style"], 4: d["food_day"],
                5: f'=SUMIFS({RPRICE},{RKEY},A{r},{RWHO},"Together")', 6: f'=COUNTIFS({RKEY},A{r},{RWHO},"Together",{RMEAL},"Yes")',
                7: o["transport_pp"], 8: o["shared"], 9: o["activities_pp"],
                10: f'=INDEX({FPRICE},MATCH("OUT|"&AC{r},{FKEY},0))', 11: f'=INDEX({FPRICE},MATCH("HER|"&AD{r},{FKEY},0))',
                12: f'=INDEX({FPRICE},MATCH("HIS|"&AE{r},{FKEY},0))',
                13: f"=2*D{r}*{DAYS}-2*{CREDIT}*D{r}*F{r}", 14: f"=2*G{r}+H{r}", 15: f"=2*I{r}", 16: f"=2*{MISC}",
                17: f"=E{r}+M{r}+N{r}+O{r}+P{r}",
                18: d["solo_food"], 19: f'=SUMIFS({RPRICE},{RKEY},A{r},{RWHO},"You")', 20: f'=COUNTIFS({RKEY},A{r},{RWHO},"You",{RMEAL},"Yes")',
                21: f"=R{r}*{DSOLO}-{CREDIT}*R{r}*T{r}", 22: o["solo_transport"], 23: o["solo_activities"], 24: f"=S{r}+U{r}+V{r}+W{r}",
                25: f"=2*J{r}+K{r}+L{r}+Q{r}+X{r}", 26: f"=J{r}+L{r}+Q{r}/2+X{r}", 27: f"=J{r}+K{r}+Q{r}/2",
                28: f'=B{r}&" | "&C{r}', 29: out_k, 30: her_k, 31: his_k,
            }
            for c, v in vals.items():
                font = BLUE if c in (4, 7, 8, 9, 18, 22, 23) else GREEN if c in (5, 6, 10, 11, 12, 19, 20) else BLACK
                put(wp, r, c, v, font, EUR if 4 <= c <= 27 and c not in (6, 20) else None)
            r += 1
    PL_LAST = r - 1
    rng = lambda col: f"Plans!${col}$2:${col}${PL_LAST}"
    put(wp, r + 1, 1, "Food = 2 × food/day × days − meal credit for nights with dinner + breakfast. Your share = your flights + half of the days together + your days alone. Her share = her flights + half of the days together.")

    # ------------------------------------------------ Scores
    wsc = wb.create_sheet("Scores")
    crit = D["criteria"]
    header(wsc, 1, ["Trip"] + [c[1] for c in crit] + ["Fit /100", "Forecast /5 (27 Oct run)"], [40] + [11] * (len(crit) + 2))
    put(wsc, 2, 1, "Weight", BOLD)
    for j, (k, _) in enumerate(crit, 2):
        put(wsc, 2, j, D["weights"][k], BLUE)
    wcols = f"$B$2:${L(len(crit) + 1)}$2"
    score_row = {}
    for i, d in enumerate(D["dests"], 3):
        score_row[d["id"]] = i
        put(wsc, i, 1, d["title"])
        for j, (k, _) in enumerate(crit, 2):
            c = put(wsc, i, j, d["scores"][k], BLUE)
            if k == "weather":
                c.comment = Comment("Half the forecast score (ECMWF run of 27 Oct, 31 Oct–7 Nov), half the November climate.", "Claude")
        fc = len(crit) + 2
        put(wsc, i, fc, f"=ROUND(100*SUMPRODUCT(B{i}:{L(len(crit) + 1)}{i},{wcols})/(5*SUM({wcols})),0)")
        put(wsc, i, fc + 1, d["fc"]["score"], BLUE)
    FITCOL, FCCOL = L(len(crit) + 2), L(len(crit) + 3)

    # ------------------------------------------------ Overview
    wo = wb.create_sheet("Overview", 1)
    header(wo, 1, ["Trip", "Plan", "Fit /100", "Forecast /5", "Places together", "Places alone", "Transport (h)", "Days together", "Your days alone",
                   "Total for two", "Your share", "Her share", "Route"], [40, 9, 9, 11, 10, 10, 11, 12, 12, 13, 12, 11, 100])
    r = 2
    for d in D["dests"]:
        for o in d["options"]:
            pr = plan_rows[f"{d['id']}|{o['style']}"]
            sr = score_row[d["id"]]
            put(wo, r, 1, d["title"]); put(wo, r, 2, o["style"])
            put(wo, r, 3, f"=Scores!{FITCOL}{sr}", GREEN); put(wo, r, 4, f"=Scores!{FCCOL}{sr}", GREEN)
            put(wo, r, 5, o["beds"]); put(wo, r, 6, o["solo_beds"]); put(wo, r, 7, round(o["travel"] / 60, 1), fmt="0.0")
            put(wo, r, 8, f"=Plans!Q{pr}", GREEN, EUR); put(wo, r, 9, f"=Plans!X{pr}", GREEN, EUR)
            put(wo, r, 10, f"=Plans!Y{pr}", GREEN, EUR); put(wo, r, 11, f"=Plans!Z{pr}", GREEN, EUR); put(wo, r, 12, f"=Plans!AA{pr}", GREEN, EUR)
            put(wo, r, 13, " → ".join(o["route"]))
            r += 1
    put(wo, r + 1, 1, "Euros. Days together = rooms, food, transport, activities and misc for the two of you (split 50/50). Transport counts every leg from Central on Sat 31 Oct to your flight home.")

    # ------------------------------------------------ My plan
    wm = wb.create_sheet("My plan", 1)
    wm.column_dimensions["A"].width = 38; wm.column_dimensions["B"].width = 46
    put(wm, 1, 1, "My plan: pick from the yellow cells", BOLD)
    trips = [d["title"] for d in D["dests"]]
    put(wm, 3, 1, "Trip"); c = put(wm, 3, 2, [d["title"] for d in D["dests"] if d["id"] == "jp-hakone"][0]); c.fill = YELLOW
    put(wm, 4, 1, "Plan"); c = put(wm, 4, 2, "Chill")
    for i, t in enumerate(trips, 1):
        put(wm, i, 26, t)
    wm.column_dimensions["Z"].hidden = True
    dv1 = DataValidation(type="list", formula1=f"=$Z$1:$Z${len(trips)}", allow_blank=False)
    dv2 = DataValidation(type="list", formula1='"Chill"', allow_blank=False)
    for dv, cell in ((dv1, "B3"), (dv2, "B4")):
        wm.add_data_validation(dv); dv.add(cell)
    put(wm, 5, 1, "Row in Plans"); put(wm, 5, 2, f'=MATCH(B3&" | "&B4,{rng("AB")},0)')
    lines = [
        ("Flights out, both of you (Sat 31 Oct)", f'=2*INDEX({rng("J")},B5)'),
        ("Her flight back to HK (Wed 4 Nov)", f'=INDEX({rng("K")},B5)'),
        ("Your flight home (Sat 7 Nov)", f'=INDEX({rng("L")},B5)'),
        ("Rooms together, 4 nights", f'=INDEX({rng("E")},B5)'),
        ("Food together", f'=INDEX({rng("M")},B5)'),
        ("Transport together", f'=INDEX({rng("N")},B5)'),
        ("Activities together", f'=INDEX({rng("O")},B5)'),
        ("Misc", f'=INDEX({rng("P")},B5)'),
        ("Your days alone (rooms, food, transport, activities)", f'=INDEX({rng("X")},B5)'),
    ]
    r0 = 7
    put(wm, r0, 1, "Cost", BOLD); put(wm, r0, 2, "€", BOLD)
    for i, (a, f) in enumerate(lines, r0 + 1):
        put(wm, i, 1, a); put(wm, i, 2, f, GREEN, EUR)
    tr = r0 + len(lines) + 1
    put(wm, tr, 1, "Total for two", BOLD); put(wm, tr, 2, f"=SUM(B{r0 + 1}:B{tr - 1})", BOLD, EUR)
    put(wm, tr + 1, 1, "Check: matches Plans total"); put(wm, tr + 1, 2, f'=IF(ABS(B{tr}-INDEX({rng("Y")},B5))<0.5,"OK","Mismatch")')
    put(wm, tr + 3, 1, "Your share", BOLD); put(wm, tr + 3, 2, f'=INDEX({rng("Z")},B5)', BOLD, EUR)
    put(wm, tr + 4, 1, "Her share", BOLD); put(wm, tr + 4, 2, f'=INDEX({rng("AA")},B5)', BOLD, EUR)
    put(wm, tr + 5, 1, "Fit /100"); put(wm, tr + 5, 2, f'=INDEX(Overview!$C$2:$C${PL_LAST},B5)', GREEN)
    put(wm, tr + 7, 1, "Each of you pays your own flights; the days together are split 50/50; your days alone are yours.")

    # ------------------------------------------------ Day by day
    wd = wb.create_sheet("Day by day")
    header(wd, 1, ["Trip", "Plan", "Day", "Date", "Who", "Title", "Pace", "Morning", "Afternoon", "Evening", "Sleep", "If it rains", "Transport (min)", "Legs"],
           [26, 8, 5, 11, 15, 30, 8, 45, 45, 35, 24, 30, 10, 50])
    r = 2
    for d in D["dests"]:
        for o in d["options"]:
            for i, x in enumerate(o["days"]):
                vals = [d["short"], o["style"], i + 1, D["dates"][i], WHO[x["who"]], x["t"], x["e"], x["am"], x["pm"], x["eve"], x["sleep"], x["b"] or "",
                        sum(l[2] for l in x["tr"]), " · ".join(f"{l[1]} ({l[2]} min)" for l in x["tr"])]
                for c, v in enumerate(vals, 1):
                    put(wd, r, c, v, wrap=c in (6, 8, 9, 10, 12, 14))
                r += 1

    # ------------------------------------------------ Checks
    wc = wb.create_sheet("Checks")
    header(wc, 1, ["Where", "What", "Status", "Finding", "Source", "Link"], [10, 34, 26, 80, 30, 50])
    order = {"out": 0, "todo": 1, "warn": 2, "ok": 3}
    for i, c in enumerate(sorted(D["checks"], key=lambda c: order[c[2]]), 2):
        for j, v in enumerate([c[0], c[1], CHK[c[2]], c[3], c[4], c[5]], 1):
            cell = put(wc, i, j, v, wrap=j == 4)
            if j == 6:
                cell.hyperlink = v; cell.font = Font(name=F, color="0563C1", underline="single")
            if j == 3:
                cell.fill = PatternFill("solid", fgColor={"ok": "DDF0E3", "warn": "FBEBD0", "out": "F6D7D5", "todo": "EEEEEE"}[c[2]])
    put(wc, len(D["checks"]) + 3, 1, f"Checked {D['checked']}.")

    # ------------------------------------------------ Sources
    wsr = wb.create_sheet("Sources")
    header(wsr, 1, ["Source", "Link", "Note"], [60, 70, 40])
    for i, (n, u, note) in enumerate(D["sources"], 2):
        put(wsr, i, 1, n); c = put(wsr, i, 2, u); c.hyperlink = u; c.font = Font(name=F, color="0563C1", underline="single"); put(wsr, i, 3, note)

    wb.save(path)
